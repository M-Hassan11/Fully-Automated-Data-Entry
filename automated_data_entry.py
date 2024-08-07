from tkinter import *
from tkinter.ttk import Combobox
from tkinter import messagebox
import openpyxl
import pathlib

def submit():
    # Data validation
    if not nameValue.get() or not contactValue.get() or not ageValue.get() or not gender_combobox.get() or not addressEntry.get("1.0", END).strip():
        messagebox.showerror("Error", "All fields must be filled out.")
        return
    if not contactValue.get().isdigit() or len(contactValue.get()) != 10:
        messagebox.showerror("Error", "Contact number must be a 10-digit number.")
        return
    if not ageValue.get().isdigit() or int(ageValue.get()) <= 0:
        messagebox.showerror("Error", "Age must be a positive integer.")
        return
    
    # Save data to Excel
    file_path = pathlib.Path("data.xlsx")
    if file_path.exists():
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Full Name", "Contact No", "Age", "Gender", "Address"])

    ws.append([nameValue.get(), contactValue.get(), ageValue.get(), gender_combobox.get(), addressEntry.get("1.0", END).strip()])
    wb.save(file_path)
    messagebox.showinfo("Success", "Data saved successfully!")
    clear_fields()

def clear_fields():
    nameEntry.delete(0, END)
    contactEntry.delete(0, END)
    ageEntry.delete(0, END)
    gender_combobox.set('Male')
    addressEntry.delete("1.0", END)

def exit_app():
    root.destroy()

def add_placeholder(widget, placeholder):
    def on_focus_in(event):
        if widget.get() == placeholder:
            widget.delete(0, 'end')  # Remove placeholder text
            widget.config(fg='black')  # Set text color to black

    def on_focus_out(event):
        if widget.get() == '':
            widget.insert(0, placeholder)  # Add placeholder text
            widget.config(fg='grey')  # Set placeholder color

    widget.insert(0, placeholder)
    widget.config(fg='grey')
    widget.bind("<FocusIn>", on_focus_in)
    widget.bind("<FocusOut>", on_focus_out)

root = Tk()
root.title("Data Entry")
root.geometry('700x415+300+200')
root.resizable(False, False)
root.configure(bg='#292626')

# Icon
icon_image = PhotoImage(file="logo.png")
root.iconphoto(False, icon_image)

# Labels and Entries
Label(root, text="Fill out this entry form", font="arial 14", bg='#292626', fg="#fff").place(x=15, y=15)
Label(root, text="Full Name", font="arial 12", bg='#292626', fg="#fff").place(x=50, y=70)
Label(root, text="Contact No", font="arial 12", bg='#292626', fg="#fff").place(x=50, y=130)
Label(root, text="Age", font="arial 12", bg='#292626', fg="#fff").place(x=50, y=190)
Label(root, text="Gender", font="arial 12", bg='#292626', fg="#fff").place(x=348, y=190)
Label(root, text="Address", font="arial 12", bg='#292626', fg="#fff").place(x=50, y=240)

nameValue = StringVar()
contactValue = StringVar()
ageValue = StringVar()

nameEntry = Entry(root, textvariable=nameValue, width=45, bd=2, font=("Arial", 13))
contactEntry = Entry(root, textvariable=contactValue, width=45, bd=2, font=("Arial", 12))
ageEntry = Entry(root, textvariable=ageValue, width=14, bd=2, font=("Arial", 12))

gender_combobox = Combobox(root, values=['Male', 'Female'], font='arial 14', state='r', width=12)
gender_combobox.place(x=426, y=190)
gender_combobox.set('Male')

addressEntry = Text(root, width=37, height=3, bd=2, font=10)

# Add placeholders
add_placeholder(nameEntry, "Enter Full Name")
add_placeholder(contactEntry, "Enter Contact No")
add_placeholder(ageEntry, "Enter Age")

nameEntry.place(x=170, y=70)
contactEntry.place(x=170, y=130)
ageEntry.place(x=170, y=190)
addressEntry.place(x=170, y=240)

Button(root, text="Submit", bg='lightblue', fg='black', width=15, height=2, command=submit).place(x=170, y=350)
Button(root, text="Clear", bg='lightblue', fg='black', width=15, height=2, command=clear_fields).place(x=320, y=350)
Button(root, text="Exit", bg='lightblue', fg='black', width=15, height=2, command=exit_app).place(x=470, y=350)
root.mainloop()

